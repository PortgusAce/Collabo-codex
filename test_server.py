import asyncio
import json
from pathlib import Path
import sys

import pytest
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

from server import LocalFiles, build_server


@pytest.fixture
def files(tmp_path):
    return LocalFiles({"roots": [str(tmp_path)], "conversion_timeout_seconds": 45})


def test_chinese_text_and_paging(files, tmp_path):
    p = tmp_path / "中文资料.txt"
    p.write_text("麻醉风险\nsecond line", encoding="utf-8")
    first = files.read(str(p), max_chars=3)
    second = files.read(str(p), offset=first["next_offset"])
    assert first["content"] + second["content"] == p.read_bytes().decode("utf-8")
    assert files.search("资料")["matches"][0]["path"] == str(p)
    assert files.search("*.TXT")["matches"][0]["path"] == str(p)
    p.write_bytes("中文编码测试".encode("gb18030"))
    assert files.read(str(p))["content"] == "中文编码测试"
    p.write_text("Unicode中文", encoding="utf-16")
    assert files.read(str(p))["content"] == "Unicode中文"


def test_listing_search_limits(files, tmp_path):
    for n in range(5):
        (tmp_path / f"file{n}.txt").write_text(str(n))
    first = files.list_directory(str(tmp_path), limit=2)
    second = files.list_directory(str(tmp_path), offset=first["next_offset"], limit=2)
    assert len(first["entries"]) == len(second["entries"]) == 2
    assert {x["name"] for x in first["entries"]}.isdisjoint(x["name"] for x in second["entries"])
    assert files.search("file", limit=2)["incomplete"]
    (tmp_path / "deep").mkdir()
    assert files.search("absent", max_depth=0)["incomplete"]


def test_path_boundary_and_secrets(files, tmp_path):
    outside = tmp_path.parent / "outside.txt"
    outside.write_text("must not read")
    for value in (str(outside), str(tmp_path / ".." / "outside.txt"), "https://example.com", "\\\\server\\share\\file.txt", "relative.txt"):
        with pytest.raises((ValueError, PermissionError)):
            files.read(value)
    for name in (".env", ".env.local", "id_rsa", "private.pem", "auth.json"):
        p = tmp_path / name
        p.write_text("secret")
        with pytest.raises(PermissionError):
            files.read(str(p))
    assert not files.list_directory(str(tmp_path))["entries"]
    assert not files.search("*")["matches"]
    private = tmp_path / ".runtime"
    private.mkdir()
    (private / "runtime-key.dpapi").write_text("synthetic encrypted credential")
    with pytest.raises(PermissionError):
        files.read(str(private / "runtime-key.dpapi"))
    assert not files.search("runtime")["matches"]


def test_junction_escape(files, tmp_path):
    import subprocess
    outside = tmp_path.parent / "outside-directory"
    outside.mkdir(exist_ok=True)
    (outside / "no.txt").write_text("outside")
    link = tmp_path / "link"
    if sys.platform == "win32":
        command = "New-Item -ItemType Junction -Path '" + str(link).replace("'", "''") + "' -Target '" + str(outside).replace("'", "''") + "'"
        result = subprocess.run(["powershell.exe", "-NoProfile", "-Command", command], capture_output=True)
        if result.returncode:
            pytest.skip("Junction creation is unavailable")
    else:
        link.symlink_to(outside, target_is_directory=True)
    with pytest.raises(PermissionError):
        files.read(str(link / "no.txt"))


def test_unsupported_large_and_bad_inputs(files, tmp_path):
    p = tmp_path / "file.doc"
    p.write_bytes(b"legacy")
    with pytest.raises(ValueError, match="Unsupported format"):
        files.read(str(p))
    p = tmp_path / "file.txt"
    p.write_bytes(b"abcdef")
    files.max_bytes = 3
    with pytest.raises(ValueError, match="exceeds"):
        files.read(str(p))
    with pytest.raises(ValueError):
        files.search("", limit=0)
    with pytest.raises(ValueError):
        files.read(str(p), offset=-1)


def test_word_powerpoint_excel(files, tmp_path):
    doc = Document()
    doc.add_heading("中文文档测试", 1)
    doc.add_paragraph("Word body marker 123")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "合成测试"
    word = tmp_path / "中文 Word.docx"
    doc.save(word)
    content = files.read(str(word))["content"]
    assert "中文文档测试" in content and "Word body marker 123" in content and "合成测试" in content
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "中文幻灯片"
    slide.placeholders[1].text = "PPT body marker 456"
    slide.notes_slide.notes_text_frame.text = "Speaker note marker 789"
    ppt = tmp_path / "测试 slides.pptx"
    presentation.save(ppt)
    content = files.read(str(ppt))["content"]
    assert "中文幻灯片" in content and "PPT body marker 456" in content and "Speaker note marker 789" in content
    workbook = Workbook()
    workbook.active.append(["synthetic", 42])
    excel = tmp_path / "测试.xlsx"
    workbook.save(excel)
    content = files.read(str(excel))["content"]
    assert "synthetic" in content and "42" in content


def test_real_mcp_stdio(tmp_path):
    config = tmp_path / "settings.json"
    config.write_text(json.dumps({"roots": [str(tmp_path)]}))
    (tmp_path / "协议测试.txt").write_text("protocol smoke test", encoding="utf-8")

    async def run():
        params = StdioServerParameters(command=sys.executable,
                                       args=[str(Path(__file__).with_name("server.py")), "--config", str(config)])
        async with stdio_client(params) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                tools = (await session.list_tools()).tools
                assert {t.name for t in tools} == {"list_roots", "list_files", "search_files", "read_file"}
                assert all(t.annotations.readOnlyHint for t in tools)
                for name, args in [("list_roots", {}), ("list_files", {"path": str(tmp_path)}),
                                   ("search_files", {"query": "协议"}),
                                   ("read_file", {"path": str(tmp_path / "协议测试.txt")})]:
                    result = await session.call_tool(name, args)
                    assert not result.isError
                assert "protocol smoke test" in result.content[0].text
                denied = await session.call_tool("read_file", {"path": str(tmp_path.parent)})
                assert denied.isError

    asyncio.run(run())
